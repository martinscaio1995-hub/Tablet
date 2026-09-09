from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

BLUE  = Font(name="Arial", size=10, color="0000FF")
BLACK = Font(name="Arial", size=10)
BOLD  = Font(name="Arial", size=10, bold=True)
TITLE = Font(name="Arial", size=14, bold=True)
SUB   = Font(name="Arial", size=10, italic=True, color="595959")
HEADF = Font(name="Arial", size=10, bold=True, color="FFFFFF")
YEL   = PatternFill("solid", fgColor="FFFF00")
HEAD  = PatternFill("solid", fgColor="1F3864")
BAND  = PatternFill("solid", fgColor="F2F2F2")
THIN  = Side(style="thin", color="BFBFBF")
BOX   = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BRL   = 'R$ #,##0.00;(R$ #,##0.00);-'
PCT   = '0.0%;(0.0%);-'
NUM   = '#,##0.0;(#,##0.0);-'

wb = Workbook()

# ---------------- Sheet 1: Custo por Peça ----------------
ws = wb.active
ws.title = "Custo por Peca"

ws["A1"] = "Leilão de Bonecos 3D — Custo e Lance Mínimo"
ws["A1"].font = TITLE
ws["A2"] = "Bambu Lab A1 · TikTok Shop. Preencha apenas as células AMARELAS (texto azul)."
ws["A2"].font = SUB

ws["A4"] = "PREMISSAS"
ws["A4"].font = BOLD

prem = [
    ("Preço do filamento (R$/kg)",            110.0,  '#,##0.00', "Preço que você paga por rolo de 1 kg de PLA."),
    ("Tarifa de energia (R$/kWh)",              0.95,  '#,##0.0000', "Veja na sua conta de luz: total da fatura / kWh consumidos."),
    ("Potência média da A1 (W)",              100.0,  '#,##0',    "Média realista com bico e mesa aquecidos em PLA. Troque se medir com wattímetro."),
    ("Embalagem + etiqueta (R$/peça)",          4.00, '#,##0.00', "Caixa/saco bolha, fita, etiqueta e brinde."),
    ("Taxa de falha / refugo (%)",              0.10, '0%',       "Fração das impressões que dá errado. 10% é conservador para PLA em A1."),
    ("Desgaste e manutenção (R$/hora)",         0.50, '#,##0.00', "Bico, placa PEI, correias e peças de reposição, rateados por hora de impressão."),
    ("Comissão do marketplace (%)",             0.08, '0%',       "Comissão do TikTok Shop. Confirme a taxa vigente da sua categoria no painel do vendedor."),
    ("Seu custo por hora (R$)",                25.00, '#,##0.00', "Quanto vale a sua hora. Entra no custo via minutos de manuseio da peça."),
    ("Multiplicador de preço-alvo (x)",         3.00, '#,##0.00', "Preço-alvo = custo x este número. 3x é o padrão de varejo para artesanal."),
]
r = 5
for label, val, fmt, note in prem:
    ws.cell(r, 1, label).font = BLACK
    c = ws.cell(r, 2, val)
    c.font = BLUE; c.fill = YEL; c.number_format = fmt; c.border = BOX
    c.comment = Comment(note, "Planilha")
    r += 1
FIRST_PREM = 5
P_FIL, P_KWH, P_W, P_EMB, P_FALHA, P_DESG, P_COM, P_HORA, P_MULT = [f"$B${FIRST_PREM+i}" for i in range(9)]

ws[f"A{r+1}"] = "PEÇAS"
ws[f"A{r+1}"].font = BOLD
ws[f"D{r+1}"] = "Preencha as 4 primeiras colunas. O resto é calculado."
ws[f"D{r+1}"].font = SUB

HR = r + 2          # header row
D0 = HR + 1         # first data row
DN = D0 + 23        # last data row (24 peças)

cols = [
    ("Peça", 26), ("Peso (g)", 10), ("Impressão (h)", 13), ("Manuseio (min)", 14),
    ("Filamento", 11), ("Energia", 10), ("Embalagem", 11), ("Desgaste", 10), ("Seu tempo", 11),
    ("Subtotal", 11), ("Refugo", 10), ("CUSTO TOTAL", 13),
    ("LANCE MÍNIMO", 14), ("Preço-alvo", 12), ("Lucro no alvo", 13), ("Margem", 9),
]
for i, (name, w) in enumerate(cols, start=1):
    c = ws.cell(HR, i, name)
    c.font = HEADF; c.fill = HEAD; c.border = BOX
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[HR].height = 30

for row in range(D0, DN + 1):
    ws.cell(row, 5,  f"=IF($B{row}=\"\",\"\",$B{row}*{P_FIL}/1000)")
    ws.cell(row, 6,  f"=IF($B{row}=\"\",\"\",$C{row}*{P_W}/1000*{P_KWH})")
    ws.cell(row, 7,  f"=IF($B{row}=\"\",\"\",{P_EMB})")
    ws.cell(row, 8,  f"=IF($B{row}=\"\",\"\",$C{row}*{P_DESG})")
    ws.cell(row, 9,  f"=IF($B{row}=\"\",\"\",$D{row}/60*{P_HORA})")
    ws.cell(row, 10, f"=IF($B{row}=\"\",\"\",SUM($E{row}:$I{row}))")
    ws.cell(row, 11, f"=IF($B{row}=\"\",\"\",$J{row}*{P_FALHA})")
    ws.cell(row, 12, f"=IF($B{row}=\"\",\"\",$J{row}+$K{row})")
    ws.cell(row, 13, f"=IFERROR(CEILING($L{row}/(1-{P_COM}),1),\"\")")
    ws.cell(row, 14, f"=IFERROR(CEILING($L{row}*{P_MULT}/(1-{P_COM}),1),\"\")")
    ws.cell(row, 15, f"=IFERROR($N{row}*(1-{P_COM})-$L{row},\"\")")
    ws.cell(row, 16, f"=IFERROR($O{row}/$N{row},\"\")")
    for col in range(1, 17):
        c = ws.cell(row, col)
        c.border = BOX
        c.font = BLUE if col <= 4 else BLACK
        if col <= 4:
            c.fill = YEL
        elif row % 2 == 0:
            c.fill = BAND
        if col in (2, 3, 4):
            c.number_format = NUM
        elif col == 16:
            c.number_format = PCT
        elif col >= 5:
            c.number_format = BRL
    for col in (12, 13):
        ws.cell(row, col).font = Font(name="Arial", size=10, bold=True)

# exemplo preenchido
for col, val in zip((1, 2, 3, 4), ("Dragão articulado M (exemplo)", 95, 5.5, 8)):
    ws.cell(D0, col, val)

ws.cell(DN + 2, 1, "LEGENDA").font = BOLD
notes = [
    "Células amarelas com texto azul = você preenche. Todo o resto é fórmula, não edite.",
    "LANCE MÍNIMO (col. M) = o valor abaixo do qual você trabalha de graça. Nunca abra lance abaixo disso.",
    "Preço-alvo (col. N) = preço justo de catálogo, já com o multiplicador e a comissão embutidos.",
    "Lucro e margem consideram a comissão do marketplace, mas NÃO o frete — lance o frete na aba Registro de Vendas.",
    "A linha de exemplo é só um modelo de formato: apague e coloque as suas peças.",
]
for i, n in enumerate(notes):
    ws.cell(DN + 3 + i, 1, ("• " + n)).font = SUB
ws.freeze_panes = f"A{D0}"
ws.column_dimensions["A"].width = 30

# ---------------- Sheet 2: Registro de Vendas ----------------
v = wb.create_sheet("Registro de Vendas")
v["A1"] = "Registro de Vendas — o que cada peça realmente fechou"
v["A1"].font = TITLE
v["A2"] = "Preencha após cada live. O custo é puxado da aba Custo por Peca pelo nome da peça (escreva igual)."
v["A2"].font = SUB

vhead = [("Data", 12), ("Peça", 28), ("Canal", 14), ("Preço final", 12),
         ("Comissão", 11), ("Custo da peça", 13), ("Frete que eu paguei", 17), ("LUCRO", 12)]
VH = 4
for i, (name, w) in enumerate(vhead, start=1):
    c = v.cell(VH, i, name)
    c.font = HEADF; c.fill = HEAD; c.border = BOX
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    v.column_dimensions[get_column_letter(i)].width = w
v.row_dimensions[VH].height = 30

V0, VN = VH + 1, VH + 60
for row in range(V0, VN + 1):
    v.cell(row, 5, f"=IF($D{row}=\"\",\"\",$D{row}*'Custo por Peca'!{P_COM})")
    v.cell(row, 6, f"=IFERROR(INDEX('Custo por Peca'!$L${D0}:$L${DN},MATCH($B{row},'Custo por Peca'!$A${D0}:$A${DN},0)),\"\")")
    v.cell(row, 8, f"=IF($D{row}=\"\",\"\",$D{row}-$E{row}-IF($F{row}=\"\",0,$F{row})-IF($G{row}=\"\",0,$G{row}))")
    for col in range(1, 9):
        c = v.cell(row, col)
        c.border = BOX
        c.font = BLUE if col in (1, 2, 3, 4, 7) else BLACK
        if col in (1, 2, 3, 4, 7):
            c.fill = YEL
        elif row % 2 == 0:
            c.fill = BAND
        if col == 1:
            c.number_format = 'DD/MM/YYYY'
        elif col >= 4:
            c.number_format = BRL
    v.cell(row, 8).font = Font(name="Arial", size=10, bold=True)

v.cell(V0, 1, "09/09/2026").font = BLUE
v.cell(V0, 2, "Dragão articulado M (exemplo)")
v.cell(V0, 3, "Live TikTok")
v.cell(V0, 4, 89)
v.cell(V0, 7, 0)

TOT = VN + 2
v.cell(TOT, 3, "TOTAIS").font = BOLD
for col, letter in ((4, "D"), (5, "E"), (6, "F"), (7, "G"), (8, "H")):
    c = v.cell(TOT, col, f"=SUM({letter}{V0}:{letter}{VN})")
    c.font = BOLD; c.number_format = BRL; c.border = BOX
v.cell(TOT + 1, 3, "Ticket médio").font = BOLD
c = v.cell(TOT + 1, 4, f"=IFERROR(D{TOT}/COUNT(D{V0}:D{VN}),0)")
c.font = BOLD; c.number_format = BRL
v.cell(TOT + 2, 3, "Margem média").font = BOLD
c = v.cell(TOT + 2, 4, f"=IFERROR(H{TOT}/D{TOT},0)")
c.font = BOLD; c.number_format = PCT

v.cell(TOT + 4, 1, "Escreva o nome da peça EXATAMENTE como está na aba Custo por Peca, senão o custo não é encontrado.").font = SUB
v.freeze_panes = f"A{V0}"

wb.save("/home/user/Tablet/leilao-3d/custo-lance-minimo.xlsx")
print("ok")
